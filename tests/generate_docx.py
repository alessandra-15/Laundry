"""
Generate WashFlow Test Cases Word Document
"""
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

doc = Document()

# ─── Styles ──────────────────────────────────────────────────────
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)

# ─── Helper functions ────────────────────────────────────────────
def set_cell_shading(cell, color):
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color)
    shading.set(qn('w:val'), 'clear')
    cell._tc.get_or_add_tcPr().append(shading)

def add_header_row(table, headers, color="0076A8"):
    row = table.rows[0]
    for i, header in enumerate(headers):
        cell = row.cells[i]
        cell.text = header
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(10)
        set_cell_shading(cell, color)

def add_test_row(table, tc_id, test_name, description, steps, expected, priority="Medium"):
    row = table.add_row()
    cells = row.cells
    cells[0].text = tc_id
    cells[1].text = test_name
    cells[2].text = description
    cells[3].text = steps
    cells[4].text = expected
    cells[5].text = priority
    for cell in cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(9)
    # Alternate row shading
    row_idx = len(table.rows) - 1
    if row_idx % 2 == 0:
        for cell in cells:
            set_cell_shading(cell, "F0F6FA")

# ─── Title Page ──────────────────────────────────────────────────
doc.add_paragraph("")
doc.add_paragraph("")
doc.add_paragraph("")

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("WashFlow")
run.bold = True
run.font.size = Pt(36)
run.font.color.rgb = RGBColor(0, 118, 168)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("Laundry Management & Booking System")
run.font.size = Pt(16)
run.font.color.rgb = RGBColor(90, 113, 132)

doc.add_paragraph("")

doc_title = doc.add_paragraph()
doc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = doc_title.add_run("Software Test Cases Document")
run.bold = True
run.font.size = Pt(24)

doc.add_paragraph("")

info = doc.add_paragraph()
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = info.add_run("Testing Tool: Appium + Android Studio Emulator\n")
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(90, 113, 132)
run = info.add_run("Framework: Python pytest + Appium-Python-Client\n")
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(90, 113, 132)
run = info.add_run("Platform: Android (Chrome Mobile Browser)")
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(90, 113, 132)

doc.add_page_break()

# ─── Table of Contents ───────────────────────────────────────────
doc.add_heading("Table of Contents", level=1)
toc_items = [
    "1. Homepage Module (TC-HP-01 to TC-HP-09)",
    "2. Customer Login Module (TC-LG-01 to TC-LG-11)",
    "3. Customer Registration Module (TC-RG-01 to TC-RG-15)",
    "4. Admin Login Module (TC-AL-01 to TC-AL-09)",
    "5. Booking Module (TC-BK-01 to TC-BK-10)",
    "6. Dashboard Module (TC-CD-01 to TC-CD-05, TC-AD-01 to TC-AD-05)",
    "7. Feedback Module (TC-FB-01 to TC-FB-09)",
    "8. Payments Module (TC-PY-01 to TC-PY-08)",
]
for item in toc_items:
    p = doc.add_paragraph(item)
    p.paragraph_format.space_after = Pt(4)

# Summary table
doc.add_paragraph("")
doc.add_heading("Test Coverage Summary", level=2)
summary_table = doc.add_table(rows=1, cols=4)
summary_table.alignment = WD_TABLE_ALIGNMENT.CENTER
add_header_row(summary_table, ["Module", "Test File", "# Test Cases", "Priority"])

summary_data = [
    ("Homepage", "test_homepage.py", "9", "High"),
    ("Customer Login", "test_login.py", "11", "Critical"),
    ("Registration", "test_register.py", "15", "Critical"),
    ("Admin Login", "test_admin_login.py", "9", "Critical"),
    ("Booking", "test_booking.py", "10", "High"),
    ("Dashboard", "test_dashboard.py", "10", "High"),
    ("Feedback", "test_feedback.py", "9", "Medium"),
    ("Payments", "test_payments.py", "8", "High"),
    ("TOTAL", "", "81", ""),
]
for data in summary_data:
    row = summary_table.add_row()
    for i, val in enumerate(data):
        row.cells[i].text = val
        for p in row.cells[i].paragraphs:
            for r in p.runs:
                r.font.size = Pt(10)
    if data[0] == "TOTAL":
        for cell in row.cells:
            set_cell_shading(cell, "0076A8")
            for p in cell.paragraphs:
                for r in p.runs:
                    r.bold = True
                    r.font.color.rgb = RGBColor(255, 255, 255)

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# MODULE 1: HOMEPAGE
# ═══════════════════════════════════════════════════════════════
doc.add_heading("1. Homepage Module (index.php)", level=1)
doc.add_paragraph("Tests for the WashFlow landing page including navigation, hero section, content sections, and call-to-action buttons.")

table = doc.add_table(rows=1, cols=6)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
add_header_row(table, ["TC ID", "Test Name", "Description", "Steps", "Expected Result", "Priority"])

add_test_row(table, "TC-HP-01", "Page Load", "Homepage loads without errors",
    "1. Navigate to index.php\n2. Check page title\n3. Verify logo visible",
    "Page loads, title is not empty, WashFlow logo is displayed", "High")

add_test_row(table, "TC-HP-02", "Hero Section Content", "Hero section shows correct heading & tagline",
    "1. Navigate to index.php\n2. Locate hero section\n3. Check heading text\n4. Check tagline visibility",
    "Heading contains 'made simple', tagline is visible", "Medium")

add_test_row(table, "TC-HP-03", "Sign In Button", "Sign In button navigates to login page",
    "1. Navigate to index.php\n2. Click 'Sign In' button in hero\n3. Verify URL",
    "Browser navigates to login.php", "High")

add_test_row(table, "TC-HP-04", "Get Started Button", "Get Started button navigates to registration",
    "1. Navigate to index.php\n2. Click 'Get Started' button\n3. Verify URL",
    "Browser navigates to register.php", "High")

add_test_row(table, "TC-HP-05", "About Section Scroll", "About Us nav link scrolls to section",
    "1. Navigate to index.php\n2. Click 'About Us' nav link\n3. Verify about section visible",
    "Page scrolls to #about section", "Medium")

add_test_row(table, "TC-HP-06", "Services Section", "Services section is present",
    "1. Navigate to index.php\n2. Locate #services element",
    "Services section exists on page", "Medium")

add_test_row(table, "TC-HP-07", "Pricing Section", "Pricing section is present",
    "1. Navigate to index.php\n2. Locate #pricing element",
    "Pricing section exists on page", "Medium")

add_test_row(table, "TC-HP-08", "FAQ Section", "FAQ section is present",
    "1. Navigate to index.php\n2. Locate #faq element",
    "FAQ section exists on page", "Medium")

add_test_row(table, "TC-HP-09", "Footer Renders", "Footer section is visible at bottom",
    "1. Navigate to index.php\n2. Scroll to bottom\n3. Check for footer element",
    "Footer element is displayed", "Low")

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# MODULE 2: CUSTOMER LOGIN
# ═══════════════════════════════════════════════════════════════
doc.add_heading("2. Customer Login Module (login.php)", level=1)
doc.add_paragraph("Tests for customer authentication including valid/invalid credentials, field validation, password toggle, remember me, and navigation links.")

table = doc.add_table(rows=1, cols=6)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
add_header_row(table, ["TC ID", "Test Name", "Description", "Steps", "Expected Result", "Priority"])

add_test_row(table, "TC-LG-01", "Page Load", "Login page loads with all form elements",
    "1. Navigate to login.php\n2. Verify form #wfLoginForm visible\n3. Check email, password, button",
    "Login form, email input, password input, and login button are all displayed", "Critical")

add_test_row(table, "TC-LG-02", "Valid Login", "Valid credentials redirect to dashboard",
    "1. Navigate to login.php\n2. Enter valid email\n3. Enter valid password\n4. Click Sign In",
    "Redirects to userdashboard.php", "Critical")

add_test_row(table, "TC-LG-03", "Logout & Re-login", "User can logout and login again",
    "1. Navigate to logout.php\n2. Navigate to login.php\n3. Enter valid credentials\n4. Click Sign In",
    "Successfully logs in again, no longer on login.php", "High")

add_test_row(table, "TC-LG-04", "Invalid Email", "Wrong email shows error message",
    "1. Navigate to login.php\n2. Enter invalid email\n3. Enter any password\n4. Click Sign In",
    "Stays on login.php, error alert (.wf-alert-danger) is displayed", "Critical")

add_test_row(table, "TC-LG-05", "Invalid Password", "Correct email + wrong password shows error",
    "1. Navigate to login.php\n2. Enter valid email\n3. Enter wrong password\n4. Click Sign In",
    "Stays on login.php", "Critical")

add_test_row(table, "TC-LG-06", "Empty Email", "Empty email prevents form submission",
    "1. Navigate to login.php\n2. Leave email empty\n3. Enter password\n4. Click Sign In",
    "Stays on login.php (client-side validation prevents submit)", "High")

add_test_row(table, "TC-LG-07", "Empty Password", "Empty password prevents form submission",
    "1. Navigate to login.php\n2. Enter email\n3. Leave password empty\n4. Click Sign In",
    "Stays on login.php", "High")

add_test_row(table, "TC-LG-08", "Password Toggle", "Password visibility toggle works",
    "1. Navigate to login.php\n2. Type password\n3. Click toggle icon\n4. Verify type=text\n5. Click again\n6. Verify type=password",
    "Password toggles between hidden (type=password) and visible (type=text)", "Medium")

add_test_row(table, "TC-LG-09", "Remember Me", "Remember me checkbox is functional",
    "1. Navigate to login.php\n2. Verify checkbox unchecked\n3. Click checkbox\n4. Verify checked",
    "Checkbox toggles between checked and unchecked states", "Low")

add_test_row(table, "TC-LG-10", "Register Link", "'Create a new account' navigates to register",
    "1. Navigate to login.php\n2. Click 'Create a new account' link",
    "Navigates to register.php", "Medium")

add_test_row(table, "TC-LG-11", "Back to Home", "'Back to Home' button navigates to homepage",
    "1. Navigate to login.php\n2. Click 'Back to Home' button",
    "Navigates to index.php", "Low")

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# MODULE 3: REGISTRATION
# ═══════════════════════════════════════════════════════════════
doc.add_heading("3. Customer Registration Module (register.php)", level=1)
doc.add_paragraph("Tests for the multi-step registration wizard including personal info (Step 1), account details (Step 2), and security/password (Step 3).")

table = doc.add_table(rows=1, cols=6)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
add_header_row(table, ["TC ID", "Test Name", "Description", "Steps", "Expected Result", "Priority"])

add_test_row(table, "TC-RG-01", "Page Load", "Register page loads with Step 1 active",
    "1. Navigate to register.php\n2. Verify form visible\n3. Check Step 1 is active",
    "Form is displayed, Step 1 has 'active' class", "Critical")

add_test_row(table, "TC-RG-02", "Step 1: Missing First Name", "First name required to proceed",
    "1. Navigate to register.php\n2. Fill all Step 1 fields except first name\n3. Click Next",
    "Remains on Step 1", "High")

add_test_row(table, "TC-RG-03", "Step 1: Invalid Email", "Invalid email prevents proceeding",
    "1. Fill Step 1 with invalid email format\n2. Click Next",
    "Remains on Step 1", "High")

add_test_row(table, "TC-RG-04", "Step 1: Valid Data", "Valid Step 1 data proceeds to Step 2",
    "1. Fill all Step 1 fields with valid data\n2. Click Next",
    "Step 2 becomes active", "Critical")

add_test_row(table, "TC-RG-05", "Phone Format", "Phone field accepts PH format",
    "1. Enter +639XXXXXXXXX in phone field\n2. Verify placeholder hint",
    "Placeholder shows PH format (09XX or +639XX)", "Medium")

add_test_row(table, "TC-RG-06", "Step 2: Address Field", "Address field visible on Step 2",
    "1. Complete Step 1\n2. Proceed to Step 2\n3. Check address field",
    "Address input field is displayed", "High")

add_test_row(table, "TC-RG-07", "Step 2: Regular Account", "Selecting Regular account type",
    "1. Go to Step 2\n2. Select 'Regular' radio button",
    "Regular radio is selected", "High")

add_test_row(table, "TC-RG-08", "Step 2: Student Upload", "Student type shows ID upload",
    "1. Go to Step 2\n2. Select 'Student' radio\n3. Check upload section",
    "Student ID upload section (#wfStudentSection) becomes visible", "High")

add_test_row(table, "TC-RG-09", "Step 3: Password Fields", "Step 3 shows password fields",
    "1. Complete Steps 1-2\n2. Proceed to Step 3\n3. Check password fields",
    "Password (#wfPassword) and confirm password (#wfConfirmPassword) are visible", "Critical")

add_test_row(table, "TC-RG-10", "Password Strength", "Strength indicator updates on typing",
    "1. Go to Step 3\n2. Type a strong password\n3. Check strength indicator",
    "Strength text (#wfPwText) is displayed and updates", "Medium")

add_test_row(table, "TC-RG-11", "Password Mismatch", "Mismatched passwords prevent submit",
    "1. Go to Step 3\n2. Enter different passwords\n3. Check terms\n4. Click submit",
    "Stays on register.php, form is not submitted", "Critical")

add_test_row(table, "TC-RG-12", "Terms Required", "Terms checkbox must be checked",
    "1. Go to Step 3\n2. Fill matching passwords\n3. Do NOT check terms\n4. Click submit",
    "Stays on register.php", "High")

add_test_row(table, "TC-RG-13", "Back Button", "Back button returns to previous step",
    "1. Go to Step 3\n2. Click 'Back' button",
    "Step 2 becomes active again", "Medium")

add_test_row(table, "TC-RG-14", "Sign In Link", "'Sign in here' navigates to login",
    "1. Navigate to register.php\n2. Click 'Sign in here' link",
    "Navigates to login.php", "Low")

add_test_row(table, "TC-RG-15", "Back to Home", "Back to Home navigates to index.php",
    "1. Navigate to register.php\n2. Click 'Back to Home'",
    "Navigates to index.php", "Low")

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# MODULE 4: ADMIN LOGIN
# ═══════════════════════════════════════════════════════════════
doc.add_heading("4. Admin Login Module (admin_login.php)", level=1)
doc.add_paragraph("Tests for admin authentication including form rendering, valid/invalid credentials, CSRF protection, and password toggle.")

table = doc.add_table(rows=1, cols=6)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
add_header_row(table, ["TC ID", "Test Name", "Description", "Steps", "Expected Result", "Priority"])

add_test_row(table, "TC-AL-01", "Page Load", "Admin login page loads with all elements",
    "1. Navigate to admin_login.php\n2. Check form, username, password, button",
    "Form #wfAdminForm, username, password inputs, and login button are visible", "Critical")

add_test_row(table, "TC-AL-02", "Admin Badge", "Admin badge is displayed",
    "1. Navigate to admin_login.php\n2. Locate .wf-admin-badge",
    "Badge with 'ADMIN' text is visible", "Low")

add_test_row(table, "TC-AL-03", "CSRF Token", "CSRF hidden field is present",
    "1. Navigate to admin_login.php\n2. Check for input[name='csrf_token']",
    "CSRF token hidden input exists in the form", "High")

add_test_row(table, "TC-AL-04", "Valid Admin Login", "Valid credentials redirect to dashboard",
    "1. Enter valid admin username\n2. Enter valid password\n3. Click Sign In",
    "Redirects to dashboard.php", "Critical")

add_test_row(table, "TC-AL-05", "Invalid Username", "Invalid username shows error",
    "1. Enter fake username\n2. Enter any password\n3. Click Sign In",
    "Stays on admin_login.php, error alert displayed", "Critical")

add_test_row(table, "TC-AL-06", "Invalid Password", "Valid username + wrong password",
    "1. Enter valid username\n2. Enter wrong password\n3. Click Sign In",
    "Stays on admin_login.php", "Critical")

add_test_row(table, "TC-AL-07", "Empty Fields", "Empty fields prevent submission",
    "1. Leave both fields empty\n2. Click Sign In",
    "Stays on admin_login.php", "High")

add_test_row(table, "TC-AL-08", "Password Toggle", "Password toggle shows/hides password",
    "1. Type password\n2. Click toggle icon\n3. Verify type=text\n4. Click again",
    "Password toggles between type=password and type=text", "Medium")

add_test_row(table, "TC-AL-09", "Back to Home", "Back to Home navigates to index.php",
    "1. Click 'Back to Home' button",
    "Navigates to index.php", "Low")

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# MODULE 5: BOOKING
# ═══════════════════════════════════════════════════════════════
doc.add_heading("5. Booking Module (book_service.php)", level=1)
doc.add_paragraph("Tests for the service booking wizard including access control, service selection, delivery options, and form fields.")

table = doc.add_table(rows=1, cols=6)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
add_header_row(table, ["TC ID", "Test Name", "Description", "Steps", "Expected Result", "Priority"])

add_test_row(table, "TC-BK-01", "Unauthenticated Redirect", "Booking requires login",
    "1. Logout\n2. Navigate to book_service.php",
    "Redirects to login.php", "Critical")

add_test_row(table, "TC-BK-02", "Authenticated Access", "Logged-in customer can access",
    "1. Login as customer\n2. Navigate to book_service.php",
    "Page loads with booking content", "Critical")

add_test_row(table, "TC-BK-03", "Services Listed", "Available services are shown",
    "1. Login\n2. Navigate to book_service.php\n3. Check page content",
    "Page contains service listings", "High")

add_test_row(table, "TC-BK-04", "Walk-in Option", "Walk-in delivery is selectable",
    "1. Navigate to booking page\n2. Select walk-in option",
    "Walk-in radio/option is selected", "High")

add_test_row(table, "TC-BK-05", "Delivery Option", "Delivery option is selectable",
    "1. Navigate to booking page\n2. Select delivery option",
    "Delivery radio/option is selected", "High")

add_test_row(table, "TC-BK-06", "Date Picker", "Drop-off date picker is present",
    "1. Navigate to booking page\n2. Locate date input field",
    "Date picker input exists", "High")

add_test_row(table, "TC-BK-07", "Time Picker", "Drop-off time picker is present",
    "1. Navigate to booking page\n2. Locate time input field",
    "Time picker input exists", "High")

add_test_row(table, "TC-BK-08", "Special Instructions", "Instructions textarea is present",
    "1. Navigate to booking page\n2. Find textarea\n3. Type instructions",
    "Text can be entered in instructions field", "Medium")

add_test_row(table, "TC-BK-09", "My Bookings Page", "Customer can view their bookings",
    "1. Login\n2. Navigate to my_bookings.php",
    "My Bookings page loads successfully", "High")

add_test_row(table, "TC-BK-10", "Dashboard Accessible", "User dashboard is accessible",
    "1. Login\n2. Navigate to userdashboard.php",
    "Dashboard page loads", "Medium")

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# MODULE 6: DASHBOARD
# ═══════════════════════════════════════════════════════════════
doc.add_heading("6. Dashboard Module", level=1)
doc.add_paragraph("Tests for both customer dashboard (userdashboard.php) and admin dashboard (dashboard.php) including access control and navigation.")

doc.add_heading("6.1 Customer Dashboard (userdashboard.php)", level=2)
table = doc.add_table(rows=1, cols=6)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
add_header_row(table, ["TC ID", "Test Name", "Description", "Steps", "Expected Result", "Priority"])

add_test_row(table, "TC-CD-01", "Unauthenticated Redirect", "Dashboard requires login",
    "1. Logout\n2. Navigate to userdashboard.php",
    "Redirects to login.php", "Critical")

add_test_row(table, "TC-CD-02", "Dashboard Loads", "Dashboard loads after login",
    "1. Login as customer\n2. Navigate to userdashboard.php",
    "Page loads with dashboard/welcome content", "Critical")

add_test_row(table, "TC-CD-03", "User Name Displayed", "Customer name shown on dashboard",
    "1. Login\n2. Check dashboard content",
    "Dashboard has substantial content", "Medium")

add_test_row(table, "TC-CD-04", "Bookings Navigation", "Can navigate to bookings",
    "1. Login\n2. Go to dashboard\n3. Check for booking links",
    "Links to booking/book_service pages exist", "High")

add_test_row(table, "TC-CD-05", "Profile Access", "Can navigate to profile page",
    "1. Login\n2. Navigate to profile.php",
    "Profile page loads", "Medium")

doc.add_heading("6.2 Admin Dashboard (dashboard.php)", level=2)
table = doc.add_table(rows=1, cols=6)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
add_header_row(table, ["TC ID", "Test Name", "Description", "Steps", "Expected Result", "Priority"])

add_test_row(table, "TC-AD-01", "Unauthenticated Redirect", "Admin dashboard requires admin login",
    "1. Logout\n2. Navigate to dashboard.php",
    "Redirects to admin_login.php", "Critical")

add_test_row(table, "TC-AD-02", "Dashboard Loads", "Admin dashboard loads after login",
    "1. Login as admin\n2. Navigate to dashboard.php",
    "Dashboard loads with admin content", "Critical")

add_test_row(table, "TC-AD-03", "Booking Management", "Admin can access booking management",
    "1. Login as admin\n2. Navigate to booking_management.php",
    "Booking management page loads", "High")

add_test_row(table, "TC-AD-04", "Inventory Access", "Admin can access inventory",
    "1. Login as admin\n2. Navigate to inventory.php",
    "Inventory page loads", "High")

add_test_row(table, "TC-AD-05", "Reports Access", "Admin can access reports",
    "1. Login as admin\n2. Navigate to reports.php",
    "Reports page loads", "High")

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# MODULE 7: FEEDBACK
# ═══════════════════════════════════════════════════════════════
doc.add_heading("7. Feedback Module", level=1)
doc.add_paragraph("Tests for customer feedback submission and admin feedback management including star ratings and response functionality.")

doc.add_heading("7.1 Customer Feedback (my_feedback.php)", level=2)
table = doc.add_table(rows=1, cols=6)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
add_header_row(table, ["TC ID", "Test Name", "Description", "Steps", "Expected Result", "Priority"])

add_test_row(table, "TC-FB-01", "Login Required", "Feedback page requires login",
    "1. Logout\n2. Navigate to my_feedback.php",
    "Redirects to login.php", "High")

add_test_row(table, "TC-FB-02", "Page Loads", "Feedback page loads for customer",
    "1. Login\n2. Navigate to my_feedback.php",
    "Page loads with feedback content", "High")

add_test_row(table, "TC-FB-03", "Form Elements", "Feedback form has rating and message",
    "1. Navigate to feedback page\n2. Check for star/rating elements",
    "Rating and message elements are present", "High")

add_test_row(table, "TC-FB-04", "Empty Feedback", "Cannot submit empty feedback",
    "1. Navigate to feedback page\n2. Click submit without input",
    "Stays on feedback page", "Medium")

add_test_row(table, "TC-FB-05", "Past Feedback", "Customer can view past feedback",
    "1. Navigate to my_feedback.php\n2. Check page content",
    "Page renders with feedback history", "Medium")

doc.add_heading("7.2 Admin Feedback Management (feedback.php)", level=2)
table = doc.add_table(rows=1, cols=6)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
add_header_row(table, ["TC ID", "Test Name", "Description", "Steps", "Expected Result", "Priority"])

add_test_row(table, "TC-FB-06", "Admin Access", "Admin can access feedback management",
    "1. Login as admin\n2. Navigate to feedback.php",
    "Feedback management page loads", "High")

add_test_row(table, "TC-FB-07", "Feedback List", "Feedback entries are displayed",
    "1. Navigate to feedback.php\n2. Check for table/list elements",
    "Feedback entries displayed in table or cards", "High")

add_test_row(table, "TC-FB-08", "Star Ratings", "Entries show star ratings",
    "1. Navigate to feedback.php\n2. Check for star icons",
    "Star rating icons (.fa-star) are present", "Medium")

add_test_row(table, "TC-FB-09", "Auth Required", "Feedback management requires admin",
    "1. Logout\n2. Navigate to feedback.php",
    "Redirects to admin_login.php or login page", "High")

doc.add_page_break()

# ═══════════════════════════════════════════════════════════════
# MODULE 8: PAYMENTS
# ═══════════════════════════════════════════════════════════════
doc.add_heading("8. Payments Module", level=1)
doc.add_paragraph("Tests for customer payment views, admin payment management, GCash payment flow, and access control.")

doc.add_heading("8.1 Customer Payments (payments.php)", level=2)
table = doc.add_table(rows=1, cols=6)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
add_header_row(table, ["TC ID", "Test Name", "Description", "Steps", "Expected Result", "Priority"])

add_test_row(table, "TC-PY-01", "Login Required", "Payment page requires authentication",
    "1. Logout\n2. Navigate to payments.php",
    "Redirects to login.php", "Critical")

add_test_row(table, "TC-PY-02", "Page Loads", "Payment page loads for customer",
    "1. Login\n2. Navigate to payments.php",
    "Page loads with payment content", "High")

add_test_row(table, "TC-PY-03", "Payment History", "Payment history is displayed",
    "1. Navigate to payments.php\n2. Check content",
    "Payment history renders", "High")

add_test_row(table, "TC-PY-04", "GCash Payment", "GCash payment page accessible",
    "1. Navigate to gcash_payment.php",
    "Page renders without errors", "Medium")

doc.add_heading("8.2 Admin Payment Management (admin_payments.php)", level=2)
table = doc.add_table(rows=1, cols=6)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
add_header_row(table, ["TC ID", "Test Name", "Description", "Steps", "Expected Result", "Priority"])

add_test_row(table, "TC-PY-05", "Admin Access", "Admin can access payment management",
    "1. Login as admin\n2. Navigate to admin_payments.php",
    "Payment management page loads", "High")

add_test_row(table, "TC-PY-06", "Payment Records", "Payment records are shown",
    "1. Navigate to admin_payments.php\n2. Check content",
    "Payment data is displayed", "High")

add_test_row(table, "TC-PY-07", "Auth Required", "Admin payments require admin login",
    "1. Logout\n2. Navigate to admin_payments.php",
    "Redirects to admin_login.php", "Critical")

add_test_row(table, "TC-PY-08", "Staff Payments", "Staff payments page accessible",
    "1. Login as admin\n2. Navigate to staff_payments.php",
    "Staff payments page loads", "Medium")

# ─── Save ────────────────────────────────────────────────────────
output_path = r"c:\Users\jhayv\OneDrive\Documents\GitHub\Laundry\WashFlow_Test_Cases.docx"
doc.save(output_path)
print(f"Document saved to: {output_path}")
